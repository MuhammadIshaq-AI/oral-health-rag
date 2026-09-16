"""Build the project overview as a .docx (imports cleanly into Google Docs, images intact)."""

from pathlib import Path

import docx
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(r"D:\repo-1-austrailia\rag\oral-health-rag")
FIG = ROOT / "docs" / "figures"
OUT = ROOT / "docs" / "RAG-Patient-Assistant-Project-Overview.docx"

ACCENT = RGBColor(0x8A, 0x6D, 0x1F)
GREY = RGBColor(0x5F, 0x63, 0x68)
INK = RGBColor(0x20, 0x21, 0x24)

doc = Document()

# ---------------------------------------------------------------- base styles
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15

for name, size, color, bold in [
    ("Title", 24, INK, True),
    ("Heading 1", 15, ACCENT, True),
    ("Heading 2", 12, INK, True),
    ("Heading 3", 10.5, INK, True),
]:
    st = doc.styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.color.rgb = color
    st.font.bold = bold
    st.paragraph_format.space_before = Pt(14 if name.startswith("Heading") else 0)
    st.paragraph_format.space_after = Pt(6)

for section in doc.sections:
    section.left_margin = section.right_margin = Inches(1.0)
    section.top_margin = section.bottom_margin = Inches(0.9)


# --------------------------------------------------------------------- helpers
def para(text="", style=None, size=None, color=None, italic=False, bold=False, align=None,
         space_after=None):
    p = doc.add_paragraph(style=style)
    if text:
        run = p.add_run(text)
        run.italic = italic
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
    if align is not None:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def rich(parts, align=None, size=None, space_after=None):
    """parts: list of (text, {bold,italic,code,color,link})."""
    p = doc.add_paragraph()
    for text, opts in parts:
        if opts.get("link"):
            add_hyperlink(p, text, opts["link"])
            continue
        run = p.add_run(text)
        run.bold = opts.get("bold", False)
        run.italic = opts.get("italic", False)
        if opts.get("code"):
            run.font.name = "Consolas"
            run.font.size = Pt(size or 9.5)
        elif size:
            run.font.size = Pt(size)
        if opts.get("color"):
            run.font.color.rgb = opts["color"]
    if align is not None:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1155CC")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(color)
    rPr.append(u)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)
    return link


def bullets(items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(4)
        if isinstance(it, str):
            p.add_run(it)
        else:  # (lead, rest)
            p.add_run(it[0]).bold = True
            p.add_run(it[1])


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9.5)
        shade(hdr[i], "EFEAE0")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            for j, seg in enumerate(val if isinstance(val, list) else [(val, {})]):
                text, opts = seg if isinstance(seg, tuple) else (seg, {})
                run = p.add_run(text)
                run.font.size = Pt(9.5)
                run.bold = opts.get("bold", False)
                if opts.get("code"):
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)
    if widths:
        for r in t.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(name, caption, width=4.8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(FIG / name), width=Inches(width))
    para(caption, size=8.5, color=GREY, italic=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)


def rule():
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "D9D2C4")
    bd.append(bottom)
    pPr.append(bd)


LIVE = "https://scope-dental-assistant.vercel.app"

# ===================================================================== content
para("Retrieval-Augmented Patient Assistant", style="Title")
para("A grounded question-answering system for a dental clinic — project overview",
     size=11.5, color=GREY, italic=True, space_after=14)

para("Every answer this system produces is built only from a fixed, licensed corpus, carries "
     "numbered citations back to its sources, and is declined rather than guessed when the corpus "
     "does not cover the question.", space_after=12)

rich([("Live application:  ", {"bold": True}), (LIVE, {"link": LIVE})], space_after=2)
rich([("Health endpoint:  ", {"bold": True}), (f"{LIVE}/api/health", {"link": f"{LIVE}/api/health"}),
      ("   →   ", {}),
      ('{"ok": true, "chunks": 245, "embedModel": "gemini-embedding-001"}', {"code": True})],
     space_after=2)
rich([("Status:  ", {"bold": True}),
      ("Deployed and serving traffic. Every figure in this document was measured against the "
       "live deployment on 16 September 2026.", {})], space_after=10)
rule()

# 1
para("1.  What the project is", style="Heading 1")
para("A general-purpose chat model will answer a patient's dental question fluently and, often "
     "enough to matter, incorrectly. It will invent an opening time, quote a price that was never "
     "published, or state a clinical claim it cannot support. For a clinic-facing assistant that "
     "failure mode is not acceptable: a wrong price is a commercial problem, and a wrong clinical "
     "instruction is a safety problem.")
para("This project is a retrieval-augmented generation (RAG) system built to remove that failure "
     "mode by construction. The language model is never asked what it knows. It is handed a small "
     "set of source excerpts retrieved for the specific question, and instructed to answer from "
     "those excerpts alone, ending every factual sentence with a citation. If the excerpts do not "
     "contain the answer, it must say exactly that and point the patient at a human.")
para("The system answers two distinct kinds of question from two distinct bodies of knowledge:")
bullets([
    ("Questions about the clinic", " — what it does, who the clinicians are, how to get in touch, "
     "where it is. Answered only from the clinic's own published information."),
    ("Everyday oral-health questions", " — brushing, bleeding gums, tooth decay, children's teeth, "
     "toothache, wisdom teeth, whitening and so on. Answered only from openly licensed patient "
     "guidance published by the NHS and the WHO."),
])
para("The two corpora are searched separately, so the five clinic passages are never buried by the "
     "240 general-guidance passages. Clinic passages are only admitted when they clear a relevance "
     "threshold, so general health questions are not padded with marketing.", space_after=4)

figure("fig1-landing.png",
       "Figure 1 — The deployed assistant. The suggested starter questions deliberately span both "
       "halves of the corpus: clinic-specific and general guidance.")

# 2
para("2.  Key figures at a glance", style="Heading 1")
table(["Measure", "Value"], [
    ["Source documents scraped", "29 pages"],
    ["Indexed passages", "245"],
    ["Passage split by publisher", "NHS 191  ·  WHO 49  ·  Clinic 5"],
    ["Corpus size", "104,498 characters (≈ 26,000 tokens)"],
    ["Passage length", "median 400 chars  ·  mean 426  ·  range 83–886"],
    ["Embedding dimensions", "768"],
    ["Shipped index size", "1.2 MB JSON (base64 Float32 vectors)"],
    ["Passages sent to the model per question", "5–7  (top 5 general + up to 2 clinic)"],
    ["Model calls per question", "2  (one embedding, one completion)"],
    ["Observed end-to-end latency", "12.4 s (measured — see Figure 3)"],
    ["Evaluation set", "54 questions — 44 in-scope, 10 out-of-scope"],
    ["Databases required at runtime", "0"],
    ["Codebase", "≈ 1,220 lines TypeScript  ·  ≈ 798 lines Python"],
], widths=[3.0, 3.5])

# 3
para("3.  How a question is answered", style="Heading 1")
para("The system separates cleanly into a build-time pipeline that runs locally in Python, and a "
     "runtime path that runs entirely inside a single serverless function. Nothing from the "
     "build-time pipeline is needed in production.")

para("Build time — run once, re-run when the knowledge base changes", style="Heading 3")
bullets([
    ("Fetch. ", "Each of the 29 source URLs is downloaded and cached to disk, so rebuilds do not "
     "re-hit the publishers' servers."),
    ("Extract. ", "Navigation, banners and boilerplate are stripped and the main body is converted "
     "to Markdown, preserving heading structure."),
    ("Chunk. ", "Pages are split first on Markdown headings, then recursively to a 900-character "
     "target with 150 characters of overlap. Heading-aware splitting is what lets a citation point "
     "at \u201cGum disease → Preventing gum disease\u201d rather than at a whole page."),
    ("Embed. ", "Every passage is embedded at 768 dimensions with the RETRIEVAL_DOCUMENT task type, "
     "L2-normalised, and written out with its title, section, URL, publisher and licence into a "
     "single JSON file that ships inside the application bundle."),
], style="List Number")

para("Runtime — per question", style="Heading 3")
bullets([
    "The browser posts the question, plus the recent turns of the conversation, to /api/chat.",
    "The question is combined with the previous question before searching, so a follow-up such as "
    "\u201chow is it treated?\u201d still retrieves the right passages. This costs nothing; the usual "
    "alternative is an extra model call to rewrite the query.",
    "The combined text is embedded with the RETRIEVAL_QUERY task type and scored by cosine "
    "similarity against all 245 pre-computed vectors, in memory. At this corpus size a full scan "
    "is faster than any index lookup and removes an entire class of infrastructure.",
    "The top 5 general passages are selected, plus up to 2 clinic passages if they clear a "
    "similarity floor.",
    "Those numbered passages and the question go to the chat model under a system prompt that "
    "forbids outside knowledge, mandates per-sentence citations, and specifies the exact refusal "
    "wording.",
    "The answer returns with its source list, each source flagged by whether the answer actually "
    "cited it. Nothing is written to storage.",
], style="List Number")

figure("fig2-cited-answer.png",
       "Figure 2 — A general-guidance answer. Every factual sentence carries a citation marker, and "
       "the source cards beneath name the exact page, the section within it, and the publisher.")

# 4
para("4.  Retrieval and generation settings", style="Heading 1")
table(["Parameter", "Value", "Why"], [
    ["Embedding model", "768-dim retrieval embeddings",
     "Asymmetric task types for passages vs. queries measurably improve retrieval over one shared encoding."],
    ["Similarity", "Cosine (dot product over normalised vectors)",
     "Vectors are normalised at build time, so the runtime comparison is a plain dot product."],
    ["General passages retrieved", "5",
     "Enough for a complete answer without diluting the prompt with weak matches."],
    ["Clinic passages retrieved", "up to 2, above a 0.6 similarity floor",
     "Guarantees clinic facts are reachable; the floor stops them surfacing on unrelated questions."],
    ["Temperature", "0.1", "The task is faithful extraction, not composition."],
    ["Reasoning effort", "disabled",
     "Reasoning tokens add latency and consume the output budget; grounded answering from supplied text does not need them."],
    ["Output cap", "1,024 tokens; prompt targets under 200 words", "Patients read short answers."],
    ["Conversation memory", "last 3 question/answer pairs",
     "Enough for follow-ups; bounded so prompt size and cost stay predictable."],
    ["Input limits", "1–1,000 characters, max 12 turns", "Basic abuse and cost control at the edge."],
    ["Timeouts and retries", "20 s embed · 45 s chat · 60 s function; retry on 429/5xx at 1.5 s then 4 s",
     "Provider rate limits and overload are usually short-lived; transient failures should not reach the patient."],
], widths=[1.5, 2.2, 2.8])

# 5
para("5.  Grounding behaviour: what it does when it does not know", style="Heading 1")
para("The interesting half of a RAG system is its behaviour on questions the corpus cannot answer. "
     "Figure 3 shows the live system asked for a treatment price the clinic has not published. "
     "Rather than estimating one, it returns a visually distinct \u201cNot covered by my sources\u201d "
     "response and routes the patient to a human.")
para("Three mechanisms produce this:")
bullets([
    ("A fixed refusal string. ", "The prompt specifies the exact sentence to emit, which makes "
     "refusals detectable in code rather than a matter of interpretation. The interface styles them "
     "differently and the evaluation harness counts them."),
    ("A hard separation of corpora. ", "Questions about the clinic may only be answered from clinic "
     "passages. The model cannot substitute general NHS guidance for a fact about this specific practice."),
    ("Citation parsing. ", "The response records which of the supplied sources the answer actually "
     "cited, so an uncited — and therefore unsupported — claim is visible rather than hidden."),
])
para("The prompt additionally forbids diagnosing individuals and forbids recommending doses beyond "
     "what a source states, and the interface carries a persistent notice that the assistant gives "
     "general information rather than a diagnosis.", space_after=4)

figure("fig3-sources-and-refusal.png",
       "Figure 3 — Top: the full source list for the previous answer, with the measured response "
       "time. Bottom: an out-of-scope question about unpublished pricing, correctly refused and "
       "redirected to the clinic.")

# 6
para("6.  Evaluation", style="Heading 1")
para("A 54-question evaluation set accompanies the system, deliberately weighted towards the ways "
     "a grounded assistant fails rather than the ways it succeeds:")
table(["Category", "Count", "Content"], [
    ["In-scope", "44", "Each tagged with the source page that ought to answer it, including 4 about the clinic itself."],
    ["Out-of-scope", "10", "Unrelated topics, competitor pricing, dosing requests, and clinic details that were never published."],
], widths=[1.2, 0.8, 4.5])
para("The harness scores three things independently, because they fail independently:")
bullets([
    ("Retrieval hit@k", " — did the tagged source page appear in the retrieved set? This isolates "
     "retrieval failure from generation failure."),
    ("Groundedness", " — a model judge labels each answer grounded, partially grounded, ungrounded "
     "or refusal. A different model family can be selected as the judge to reduce self-preference bias."),
    ("Citation rate and refusal rate", " — the proportion of answers carrying citations, and the "
     "proportion of the 10 out-of-scope questions correctly declined."),
])
para("The harness can be pointed at the deployed URL, in which case it measures exactly what "
     "patients receive — production prompt, production retrieval, production model — rather than a "
     "local approximation. Runs checkpoint to disk and resume, because a full pass needs roughly "
     "160 model calls and will cross a free-tier daily quota.")

# 7
para("7.  Architecture and deployment", style="Heading 1")
para("The entire runtime is one Next.js application on Vercel. There is no separate backend "
     "service, no vector database, no cache tier and no persistent storage. The index is a build "
     "artefact committed to the repository, so a deployment is atomic: the code and the exact "
     "knowledge base it was tested against ship together, and a rollback restores both.")
table(["Layer", "Choice"], [
    ["Application", "Next.js 16 (App Router), React 19, Tailwind CSS v4, Node 22"],
    ["API", "Two serverless routes: /api/chat and /api/health"],
    ["Vector search", "In-process scan of the bundled index — no database"],
    ["Region", "Mumbai (bom1), the deployment region closest to the clinic's patients"],
    ["Build-time tooling", "Python: HTTP fetching, boilerplate-stripping extraction, heading-aware chunking"],
    ["Secrets", "API key held server-side only; never reaches the browser"],
    ["Hardening", "Security headers on every route, no-store on all API responses, framework fingerprint header removed"],
    ["CI behaviour", "Builds are skipped for commits that do not touch the application directory"],
], widths=[1.5, 5.0])
para("Provider errors are translated at the boundary: the raw upstream error is logged for "
     "debugging, while the patient sees a plain-language message. A rate-limit response, for "
     "example, becomes an apology with the clinic's contact details rather than an HTTP status code.")

# 8
para("8.  How to use the deployed application", style="Heading 1")
bullets([
    ("Open the app. ", f"Go to {LIVE} in any modern browser. There is no sign-up, no login and no "
     "cookie banner — the assistant stores nothing."),
    ("Ask a question. ", "Click one of the six suggested cards, or type into the box at the bottom "
     "and press Enter. Questions are limited to 1,000 characters."),
    ("Wait for the answer. ", "A response typically takes several seconds; the measured example in "
     "Figure 3 took 12.4 seconds. Two model calls happen in that window — embedding the question, "
     "then generating the grounded answer."),
    ("Read the citations. ", "Every factual sentence ends with one or more numbered markers. The "
     "Sources grid underneath maps each number to the page title, the specific section within it, "
     "and the publisher. The footer reports how many sources were cited and how long it took."),
    ("Verify anything you doubt. ", "The source cards identify the original publisher, so any claim "
     "can be traced back to NHS, WHO or clinic material."),
    ("Ask a follow-up. ", "Short follow-ups such as \u201chow is it treated?\u201d work, because the "
     "previous question is folded into the search. The last three exchanges are kept for context."),
    ("Expect refusals, and trust them. ", "A \u201cNot covered by my sources\u201d panel (Figure 3) "
     "means the answer genuinely is not in the corpus — unpublished prices, unpublished opening "
     "hours, or anything outside oral health. Treat this as the system working, and follow the "
     "contact route it offers."),
    ("Copy and reset. ", "Use Copy on any answer to take the text, and New chat to clear the "
     "conversation. History lives only in the browser tab and disappears when it closes."),
    ("Check availability. ", "The /api/health endpoint reports whether the service is configured "
     "and how many passages are loaded."),
], style="List Number")
rich([("If an answer does not arrive: ", {"bold": True}),
      ("a message about a usage limit means the daily model quota is exhausted — it resets, and "
       "enabling billing on the provider account removes the ceiling. A message about the model "
       "being busy is transient; the request has already been retried twice automatically, and "
       "trying again usually succeeds.", {})])

# 9
para("9.  What it is good for, and what it is not", style="Heading 1")
rich([("Strengths. ", {"bold": True}),
      ("Answers are traceable to a named publisher and section. The knowledge base is versioned "
       "alongside the code. Operating cost is two model calls per question with no infrastructure "
       "underneath. And the system degrades honestly — it declines instead of inventing. Updating "
       "the knowledge base means editing a Markdown file or a URL list and re-running two scripts.", {})])
rich([("Limitations, stated plainly.", {"bold": True})], space_after=4)
bullets([
    "It provides general information. It is not medical advice and has not been clinically validated.",
    "Clinic-specific answers are bounded by what the clinic has published. Details that do not exist "
    "in the corpus — prices, opening hours, a service list — cannot be answered until they are added, "
    "and the system will keep refusing them until they are.",
    "The general-guidance corpus reflects the day it was scraped. It does not update itself; a "
    "rebuild is required.",
    "WHO material is licensed for non-commercial use, which constrains commercial deployment. Every "
    "passage carries its licence, so this stays checkable rather than assumed.",
    "On a free provider tier the daily request quota is small. A public launch needs billing enabled.",
    "A flat scan over 245 passages is the right choice at this size and the wrong choice at 100,000. "
    "Growth past a few thousand passages is the point to introduce a real vector index.",
])

# 10
para("10.  Repository layout", style="Heading 1")
table(["Path", "Contents"], [
    ["web/", "The deployed Next.js application: interface, chat route, health route"],
    ["web/src/lib/rag/", "Retrieval, prompt and model client used at runtime"],
    ["web/src/data/knowledge.json", "The 245 pre-embedded passages (generated artefact, 1.2 MB)"],
    ["data/clinic/", "The clinic's own information, as editable Markdown"],
    ["rag/", "Python: source list, scraping, chunking, and the local pipeline used by the evaluation"],
    ["scripts/", "Two build steps: scrape & chunk, then embed & export"],
    ["eval/", "The 54-question set and the evaluation runner"],
], widths=[2.0, 4.5])

rule()
para("Source material is used under the UK Open Government Licence v3.0 (NHS) and CC BY-NC-SA 3.0 "
     "IGO (WHO); every indexed passage retains its publisher, URL and licence. All figures were "
     "measured against the live deployment on 16 September 2026.",
     size=8.5, color=GREY, italic=True)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(f"wrote {OUT}  ({OUT.stat().st_size / 1024:.0f} KB)")
